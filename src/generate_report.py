import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


INPUT_PATH = "data/reconciliation_results.csv"
OUTPUT_PATH = "output/reconciliation_report.xlsx"


def format_sheet(worksheet):

    # Header formatting
    for cell in worksheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="1F4E78"
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    # Freeze header row
    worksheet.freeze_panes = "A2"

    # Auto-size columns
    for column_cells in worksheet.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            35
        )


def generate_report():

    df = pd.read_csv(
        INPUT_PATH
    )

    total_records = len(df)

    matched_records = (
        df["status"] == "Matched"
    ).sum()

    exception_records = (
        df["status"] != "Matched"
    ).sum()

    review_records = (
        df["review_required"] == "Yes"
    ).sum()

    match_rate = (
        matched_records
        / total_records
    ) * 100

    # ---------------------------------------------------------
    # SUMMARY
    # ---------------------------------------------------------

    summary = pd.DataFrame(
        {
            "Metric": [
                "Total Reconciliation Records",
                "Matched Records",
                "Auto-Match Rate",
                "Reconciliation Findings",
                "Exceptions Requiring Review"
            ],
            "Value": [
                total_records,
                matched_records,
                round(match_rate, 2),
                exception_records,
                review_records
            ]
        }
    )

    # ---------------------------------------------------------
    # EXCEPTION BREAKDOWN
    # ---------------------------------------------------------

    exception_breakdown = (
        df[
            df["status"] != "Matched"
        ]["status"]
        .value_counts()
        .reset_index()
    )

    exception_breakdown.columns = [
        "Issue Type",
        "Count"
    ]

    # ---------------------------------------------------------
    # SEVERITY BREAKDOWN
    # ---------------------------------------------------------

    review_breakdown = (
        df[
            df["review_required"] == "Yes"
        ]["severity"]
        .value_counts()
        .reset_index()
    )

    review_breakdown.columns = [
        "Severity",
        "Count"
    ]

    # ---------------------------------------------------------
    # REVIEW TRACKER
    # ---------------------------------------------------------

    exceptions_tracker = df[
        df["review_required"] == "Yes"
    ].copy()

    exceptions_tracker = exceptions_tracker[
        [
            "transaction_id",
            "status",
            "ledger_date",
            "bank_date",
            "ledger_vendor",
            "bank_vendor",
            "ledger_amount",
            "bank_amount",
            "amount_difference",
            "date_difference_days",
            "severity",
            "review_reason"
        ]
    ]

    exceptions_tracker.columns = [
        "Transaction ID",
        "Issue Type",
        "Ledger Date",
        "Bank Date",
        "Ledger Vendor",
        "Bank Vendor",
        "Ledger Amount",
        "Bank Amount",
        "Amount Difference",
        "Date Difference (Days)",
        "Severity",
        "Review Reason"
    ]

    # ---------------------------------------------------------
    # DATA QUALITY SUMMARY
    # ---------------------------------------------------------

    data_quality = pd.DataFrame(
        {
            "Validation": [
                "Duplicate Transaction IDs",
                "Missing Vendor Fields",
                "Missing in Bank",
                "Missing in Ledger",
                "Amount Mismatches",
                "Date Mismatches"
            ],
            "Count": [
                (
                    df["status"] == "Duplicate"
                ).sum(),

                (
                    df["status"]
                    == "Missing Vendor Field"
                ).sum(),

                (
                    df["status"]
                    == "Missing in Bank"
                ).sum(),

                (
                    df["status"]
                    == "Missing in Ledger"
                ).sum(),

                (
                    df["status"]
                    == "Amount Mismatch"
                ).sum(),

                (
                    df["status"]
                    == "Date Mismatch"
                ).sum()
            ]
        }
    )

    # ---------------------------------------------------------
    # WRITE EXCEL FILE
    # ---------------------------------------------------------

    with pd.ExcelWriter(
        OUTPUT_PATH,
        engine="openpyxl"
    ) as writer:

        summary.to_excel(
            writer,
            sheet_name="Summary",
            index=False
        )

        exception_breakdown.to_excel(
            writer,
            sheet_name="Exception Breakdown",
            index=False
        )

        review_breakdown.to_excel(
            writer,
            sheet_name="Severity Breakdown",
            index=False
        )

        exceptions_tracker.to_excel(
            writer,
            sheet_name="Exceptions Tracker",
            index=False
        )

        data_quality.to_excel(
            writer,
            sheet_name="Data Quality",
            index=False
        )

    # ---------------------------------------------------------
    # FORMAT WORKBOOK
    # ---------------------------------------------------------

    workbook = load_workbook(
        OUTPUT_PATH
    )

    for worksheet in workbook.worksheets:

        format_sheet(
            worksheet
        )

    # Percentage formatting
    summary_sheet = workbook[
        "Summary"
    ]

    summary_sheet["B4"].number_format = "0.00%"

    # Since Auto-Match Rate is stored as 91 rather than
    # 0.91, display it correctly as 91.00%.
    summary_sheet["B4"] = match_rate / 100

    workbook.save(
        OUTPUT_PATH
    )

    print("=" * 65)
    print("EXCEL REPORT GENERATED")
    print("=" * 65)

    print(
        f"Total records: {total_records}"
    )

    print(
        f"Matched records: {matched_records}"
    )

    print(
        f"Auto-match rate: {match_rate:.2f}%"
    )

    print(
        f"Reconciliation findings: "
        f"{exception_records}"
    )

    print(
        f"Exceptions requiring review: "
        f"{review_records}"
    )

    print(
        f"\nExcel report saved to: "
        f"{OUTPUT_PATH}"
    )

    print("=" * 65)


if __name__ == "__main__":
    generate_report()